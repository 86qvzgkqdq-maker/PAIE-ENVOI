import pytest
from openpyxl import Workbook

from paie.clients import ErreurColonnes, charger_clients


def test_chargement_et_colonnes_accentuees(excel_clients):
    clients = charger_clients(excel_clients)
    assert len(clients) == 3  # la ligne sans email est ignorée
    assert clients[0].nom == "Dupé"
    assert clients[0].prenom == "Jean-François"
    assert clients[0].email == "jf.dupe@exemple.fr"


def test_ordre_colonnes_indifferent(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Adresse mail", "PRENOM", "NOM DE FAMILLE"])
    ws.append(["a@b.fr", "Léa", "Petit"])
    chemin = tmp_path / "c.xlsx"
    wb.save(chemin)
    (client,) = charger_clients(chemin)
    assert (client.nom, client.prenom, client.email) == ("Petit", "Léa", "a@b.fr")


def test_entetes_en_ligne_2_avec_titre_et_nom_complet(tmp_path):
    """Cas réel : ligne de titre au-dessus, colonne « Nom complet » à côté."""
    wb = Workbook()
    ws = wb.active
    ws.append(["liste_coop"])
    ws.append(["Nom complet", "Prénom", "Nom", "Email"])
    ws.append(["Thierry PAULCAN", "Thierry", "PAULCAN", "t.paulcan@exemple.fr"])
    chemin = tmp_path / "c.xlsx"
    wb.save(chemin)
    (client,) = charger_clients(chemin)
    assert (client.nom, client.prenom) == ("PAULCAN", "Thierry")


def test_colonne_manquante(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nom", "Email"])
    chemin = tmp_path / "c.xlsx"
    wb.save(chemin)
    with pytest.raises(ErreurColonnes, match="introuvable"):
        charger_clients(chemin)
